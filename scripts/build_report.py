"""Render a verified public-professional evidence brief; no network access."""
import argparse
import json
import re
from pathlib import Path
from datetime import date
from urllib.parse import urlparse
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
import networkx as nx
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def hyperlink(p, text, url):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = text
    r.append(t); h.append(r); p._p.append(h)


def render(data, out):
    out = Path(out); out.mkdir(parents=True, exist_ok=True)
    temp = out / '_render'; temp.mkdir(exist_ok=True)
    sources = data.get('sources', [])
    ids = [s['id'] for s in sources]
    if len(ids) != len(set(ids)): raise ValueError('Duplicate source IDs')
    byid = {s['id']: s for s in sources}
    def refs(row):
        result = row.get('sources', [])
        for sid in result:
            if sid not in byid: raise ValueError('Unknown source: ' + sid)
            s = byid[sid]
            u = urlparse(s.get('url', ''))
            if not s.get('verified') or u.scheme not in ('https', 'http') or not u.netloc:
                raise ValueError('Unverified or invalid source: ' + sid)
        return ' '.join('[' + s + ']' for s in result)
    for key in ('basic', 'appointments', 'biography', 'publications', 'collaborators', 'directions'):
        for row in data.get(key, []): refs(row)
    for key, field in [('basic', 'value'), ('appointments', 'organization'), ('biography', 'event'), ('directions', 'topic')]:
        for row in data.get(key, []):
            if row.get(field) and not row.get('sources'):
                raise ValueError('Nonempty factual row requires sources: ' + key)
    m = data['meta']
    if m.get('as_of') and m['as_of'] != 'YYYY-MM-DD': date.fromisoformat(m['as_of'])
    pubs = [r for r in data.get('publications', []) if r.get('title')]
    for p in pubs:
        if not p.get('sources'): raise ValueError('Publication requires evidence')
        pd = date.fromisoformat(p['date'])
        if not (date.fromisoformat(m['window_start']) <= pd <= date.fromisoformat(m['as_of'])):
            raise ValueError('Publication outside five-year window')
    if len(pubs) > 5: raise ValueError('Select at most five representative publications')
    people = [r for r in data.get('collaborators', []) if r.get('name')]
    seen = set()
    allowed = {'共同论文', '共同试验', '共同项目'}
    for p in people:
        if not p.get('id') or p['id'] == 'target' or p['id'] in seen: raise ValueError('Missing or duplicate collaborator ID')
        seen.add(p['id'])
        if p.get('group') not in ('china', 'outside_china'): raise ValueError('Unresolved institution group')
        if p.get('relationship_type') not in allowed or not p.get('sources') or not p.get('evidence'):
            raise ValueError('Direct verified collaboration evidence required')
    for group in ('china', 'outside_china'):
        if sum(p['group'] == group for p in people) > 10: raise ValueError('At most ten collaborators per group')
    doc = Document()
    sec = doc.sections[0]; sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.8)
    normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.add_heading('医生公开学术信息简报', 0)
    doc.add_paragraph(' / '.join(m.get(k, '') or '待补充' for k in ('name', 'hospital', 'department')))
    doc.add_paragraph('截止日期：%s | 五年窗口起点：%s | 版本：%s' % (m.get('as_of',''),m.get('window_start',''),m.get('version','v1')))
    doc.add_paragraph('仅整理公开职业信息。合作记录不代表私人关系；本报告不包含性格测评或个人绩效评价。')
    doc.add_heading('1. 摘要', 1); doc.add_paragraph(m.get('summary') or '未提供已核验摘要。')
    def table(title, headers, rows):
        doc.add_heading(title, 1)
        if not rows: doc.add_paragraph('未检索到足够的已核验信息。'); return
        t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Light Shading Accent 1'
        for c,h in zip(t.rows[0].cells,headers): c.text=h
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,v in zip(t.add_row().cells,row): c.text=str(v or '未公开/未核实')
    table('2. 基本信息',['字段','内容','来源'],[[r.get('field'),r.get('value'),refs(r)] for r in data.get('basic',[]) if r.get('value')])
    table('3. 协会任职',['组织与职务','任期/状态','来源'],[[r.get('organization','')+' / '+r.get('role',''),r.get('term','')+' / '+r.get('status',''),refs(r)] for r in data.get('appointments',[]) if r.get('organization')])
    table('4. 职业生平',['时间','事件','来源'],[[r.get('date'),r.get('event'),refs(r)] for r in data.get('biography',[]) if r.get('event')])
    doc.add_heading('5. 近五年代表论文',1)
    doc.add_paragraph('已核验 %d/5 篇；不足不以同名作者或窗口外论文补足。'%len(pubs))
    for i,p in enumerate(pubs,1):
        doc.add_heading('%d. %s'%(i,p['title']),2)
        doc.add_paragraph(' | '.join(p.get(k,'') for k in ('journal','date','type','author_role','identifier')))
        doc.add_paragraph('入选理由：'+p.get('selection_reason','未提供')+' '+refs(p))
    doc.add_heading('6. 公开学术合作网络',1)
    doc.add_paragraph('按合作证据时机构所在地分组，不代表国籍；连线仅表示有证据的直接合作，不表示亲密程度。')
    if people:
        available={f.name for f in font_manager.fontManager.ttflist}
        fonts=[f for f in ('Noto Sans CJK SC','Noto Sans CJK JP','Microsoft YaHei','SimHei','WenQuanYi Zen Hei') if f in available]
        if fonts: plt.rcParams['font.family']=fonts
        g=nx.Graph(); g.add_node('target')
        ordered=sorted(people,key=lambda p:(p['group'],p['name']))
        labels={'target':m.get('name') or '目标医生'}
        for i,p in enumerate(ordered,1):
            g.add_edge('target',p['id']); labels[p['id']]=str(i)+'. '+p['name']
        pos=nx.shell_layout(g,nlist=[['target'],[p['id'] for p in ordered]])
        fig,ax=plt.subplots(figsize=(12,10))
        nx.draw_networkx(g,pos,labels=labels,ax=ax,font_family=fonts[0] if fonts else 'sans-serif',font_size=9,node_size=900)
        ax.set_title('公开学术合作网络（非私人关系图）',fontfamily=fonts[0] if fonts else 'sans-serif'); ax.axis('off')
        fig.tight_layout(); image=temp/'network.png'; fig.savefig(image,dpi=180,bbox_inches='tight'); plt.close(fig)
        doc.add_picture(str(image),width=Inches(6.5))
        if not fonts: doc.add_paragraph('注意：绘图环境未检测到常见中文字体，请检查图中文字显示。')
    else: doc.add_paragraph('无已核验直接合作记录，未绘制网络图。')
    for group,label in [('china','中国机构组'),('outside_china','中国之外机构组')]:
        subset=sorted([p for p in people if p['group']==group],key=lambda p:p['name'])
        table(label+'（%d/10）'%len(subset),['姓名/机构','合作证据/日期','来源'],[[p['name']+' / '+p.get('institution',''),p['relationship_type']+'：'+p['evidence']+' / '+p.get('date',''),refs(p)] for p in subset])
    doc.add_heading('7. 学术沟通准备',1)
    doc.add_paragraph('以下为交流准备卡，不是四色性格分类，也不代表本人心理或沟通偏好。')
    cards=doc.add_table(rows=2,cols=2); cards.style='Light Shading Accent 1'
    for cell,key in zip([c for row in cards.rows for c in row.cells],('研究主题','证据问题','交流材料','待确认事项')):
        items=data.get('communication',{}).get(key,[])
        cell.text=key+'\n'+('\n'.join(str(x) for x in items) if items else '待补充；不推断个人偏好。')
    table('8. 研究方向与建议讨论话题',['主题/性质','依据与局限','来源'],[[r.get('topic','')+' / '+r.get('kind',''),r.get('basis','')+'；'+r.get('limitation',''),refs(r)] for r in data.get('directions',[]) if r.get('topic')])
    doc.add_heading('9. 来源台账',1)
    for s in sources:
        if not s.get('title'): continue
        doc.add_paragraph('[%s] %s | %s | 发布：%s | 访问：%s | %s'%(s['id'],s['title'],s.get('publisher',''),s.get('published','未知'),s.get('accessed','未知'),'已核验' if s.get('verified') else '待核验'))
        u=s.get('url',''); parsed=urlparse(u)
        if parsed.scheme in ('http','https') and parsed.netloc: hyperlink(doc.add_paragraph(),u,u)
    doc.add_heading('10. 检索覆盖与局限',1)
    for x in data.get('limitations',[]) or ['尚未记录检索覆盖与限制；交付前必须补充。']: doc.add_paragraph(str(x))
    doc.add_heading('11. 版本与修改记录',1)
    for x in data.get('changes',[]) or ['初版；尚无用户修订记录。']: doc.add_paragraph(str(x))
    doc.add_paragraph('视觉核验状态：本脚本未执行逐页渲染检查，交付前请在宿主中完成并更新本行。')
    safe=lambda s: re.sub(r'[^\w\-]+','_',str(s))[:80]
    path=out/('physician_brief_%s_%s_%s.docx'%(safe(m.get('name') or 'template'),safe(m.get('as_of') or 'undated'),safe(m.get('version') or 'v1')))
    doc.save(path)
    return path

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('evidence'); parser.add_argument('--out',required=True)
    args=parser.parse_args()
    with open(args.evidence,encoding='utf-8') as f: data=json.load(f)
    print(render(data,args.out))
